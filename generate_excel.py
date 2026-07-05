#!/usr/bin/env python3
"""Generate Excel vocabulary table and daily practice plan from extracted vocabulary data."""

import json
import os
import math
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

def load_vocabulary():
    """Load vocabulary data from JSON file."""
    with open(os.path.join(SCRIPT_DIR, 'vocabulary_data.json'), 'r', encoding='utf-8') as f:
        return json.load(f)

def create_excel(vocab_data):
    """Create Excel file with A-Z sorted vocabulary, blue header, yellow fill, frozen panes."""
    wb = Workbook()

    # === Sheet 1: 词汇总表 (A-Z) ===
    ws1 = wb.active
    ws1.title = "词汇总表A-Z"

    # Styles
    header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='微软雅黑', size=11)
    data_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    data_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Headers
    headers = ['序号', '首字母', '单词', '音标', '词性', '中文释义', '来源教材']
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Data rows
    for idx, entry in enumerate(vocab_data, 1):
        row = idx + 1
        word = entry['word']
        first_letter = word[0].upper() if word else ''

        values = [
            idx,
            first_letter,
            word,
            entry.get('phonetic', ''),
            entry.get('pos', ''),
            entry.get('meaning', ''),
            entry.get('source', '')
        ]

        for col, value in enumerate(values, 1):
            cell = ws1.cell(row=row, column=col, value=value)
            cell.font = data_font
            cell.fill = data_fill
            cell.border = thin_border
            if col in [1, 2, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = data_align

    # Column widths
    col_widths = [6, 8, 25, 25, 12, 45, 20]
    for col, width in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(col)].width = width

    # Freeze panes (freeze header row and first 3 columns)
    ws1.freeze_panes = 'D2'

    # === Sheet 2: 按首字母统计 ===
    ws2 = wb.create_sheet("按字母统计")
    for col, header in enumerate(['首字母', '单词数', '占比'], 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Count by letter
    letter_counts = {}
    for entry in vocab_data:
        letter = entry['word'][0].upper() if entry['word'] else '#'
        letter_counts[letter] = letter_counts.get(letter, 0) + 1

    total = len(vocab_data)
    for idx, (letter, count) in enumerate(sorted(letter_counts.items()), 1):
        row = idx + 1
        ws2.cell(row=row, column=1, value=letter).font = data_font
        ws2.cell(row=row, column=2, value=count).font = data_font
        ws2.cell(row=row, column=3, value=f"{count/total*100:.1f}%").font = data_font
        for col in range(1, 4):
            ws2.cell(row=row, column=col).fill = data_fill
            ws2.cell(row=row, column=col).border = thin_border
            ws2.cell(row=row, column=col).alignment = center_align

    ws2.column_dimensions['A'].width = 10
    ws2.column_dimensions['B'].width = 10
    ws2.column_dimensions['C'].width = 10
    ws2.freeze_panes = 'A2'

    # === Sheet 3: 按教材统计 ===
    ws3 = wb.create_sheet("按教材统计")
    for col, header in enumerate(['教材', '单词数', '占比'], 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    source_counts = {}
    for entry in vocab_data:
        for s in entry.get('source', '').split(', '):
            s = s.strip()
            if s:
                source_counts[s] = source_counts.get(s, 0) + 1

    for idx, (source, count) in enumerate(sorted(source_counts.items()), 1):
        row = idx + 1
        ws3.cell(row=row, column=1, value=source).font = data_font
        ws3.cell(row=row, column=2, value=count).font = data_font
        ws3.cell(row=row, column=3, value=f"{count/total*100:.1f}%").font = data_font
        for col in range(1, 4):
            ws3.cell(row=row, column=col).fill = data_fill
            ws3.cell(row=row, column=col).border = thin_border
            ws3.cell(row=row, column=col).alignment = center_align

    ws3.column_dimensions['A'].width = 15
    ws3.column_dimensions['B'].width = 10
    ws3.column_dimensions['C'].width = 10
    ws3.freeze_panes = 'A2'

    # === Sheet 4: 每日练习计划 ===
    ws4 = wb.create_sheet("每日练习计划")

    plan_headers = ['天数', '日期', '练习主题', '新词范围', '新词数', '复习词范围', '复习词数', '练习类型', '预估时长(分钟)']
    for col, header in enumerate(plan_headers, 1):
        cell = ws4.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Design daily plan: 20 new words per day + 10 review words
    words_per_day = 20
    review_per_day = 10
    total_days = math.ceil(total / words_per_day)
    start_date = datetime(2025, 7, 4)  # Start tomorrow

    for day in range(1, total_days + 1):
        row = day + 1
        date = start_date + timedelta(days=day - 1)

        # New words range
        start_idx = (day - 1) * words_per_day
        end_idx = min(start_idx + words_per_day, total)
        new_count = end_idx - start_idx

        # Review words range (from 2 days ago)
        if day >= 3:
            review_start = max(0, (day - 3) * words_per_day)
            review_end = min(review_start + review_per_day, start_idx)
            review_count = review_end - review_start
            review_range = f"第{review_start+1}-{review_end}词"
        else:
            review_count = 0
            review_range = "无"

        # Alternate practice types
        if day % 3 == 1:
            practice_type = "英译中默写 + 听写"
        elif day % 3 == 2:
            practice_type = "中译英默写 + 听写"
        else:
            practice_type = "混合默写 + 听写"

        # Time estimate
        time_est = 30

        values = [
            f"第{day}天",
            date.strftime('%Y-%m-%d'),
            f"{vocab_data[start_idx]['word'][0].upper()}开头的单词",
            f"第{start_idx+1}-{end_idx}词",
            new_count,
            review_range,
            review_count,
            practice_type,
            time_est
        ]

        for col, value in enumerate(values, 1):
            cell = ws4.cell(row=row, column=col, value=value)
            cell.font = data_font
            cell.fill = data_fill
            cell.border = thin_border
            cell.alignment = center_align

    plan_col_widths = [8, 14, 20, 16, 8, 16, 8, 22, 14]
    for col, width in enumerate(plan_col_widths, 1):
        ws4.column_dimensions[get_column_letter(col)].width = width
    ws4.freeze_panes = 'A2'

    # Save
    output_path = os.path.join(SCRIPT_DIR, '英语教材词汇表A-Z.xlsx')
    wb.save(output_path)
    print(f"Excel saved to: {output_path}")
    print(f"Total words: {total}")
    print(f"Total days: {total_days}")
    return output_path, total_days, words_per_day

if __name__ == '__main__':
    vocab_data = load_vocabulary()
    print(f"Loaded {len(vocab_data)} vocabulary entries")

    excel_path, total_days, words_per_day = create_excel(vocab_data)
    print(f"\nExcel file created: {excel_path}")
    print(f"Daily plan: {total_days} days, {words_per_day} words/day")
